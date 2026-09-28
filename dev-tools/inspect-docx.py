"""
Inspecciona el contenido de un documento Word
"""
from pathlib import Path
from docx import Document

docx_path = Path(r"c:\Github\OCITest\Descargables\Terraform\004\source\HASHICORP TERRAFORM - 004.docx")

if not docx_path.exists():
    print(f"❌ FALTA: {docx_path}")
    exit(1)

doc = Document(docx_path)

print(f"📄 Documento: {docx_path.name}")
print(f"📊 Párrafos totales: {len(doc.paragraphs)}")
print(f"\n{'='*80}")
print("PRIMEROS 50 PÁRRAFOS (para entender el formato):")
print(f"{'='*80}\n")

for i, para in enumerate(doc.paragraphs[:50], 1):
    text = para.text.strip()
    if text:
        # Mostrar con prefijo para identificar patrones
        if text.lower().startswith('question'):
            print(f"[Q] {i:3}: {text[:100]}")
        elif text.lower().startswith('answer'):
            print(f"[A] {i:3}: {text[:100]}")
        elif text.lower().startswith('correct'):
            print(f"[C] {i:3}: {text[:100]}")
        else:
            print(f"    {i:3}: {text[:100]}")

print(f"\n{'='*80}")
print("\nÚLTIMOS PÁRRAFOS (para ver final):")
print(f"{'='*80}\n")

for i, para in enumerate(doc.paragraphs[-10:], len(doc.paragraphs)-9):
    text = para.text.strip()
    if text:
        print(f"{i:3}: {text[:100]}")
